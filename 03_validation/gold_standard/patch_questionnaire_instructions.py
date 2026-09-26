# -*- coding: utf-8 -*-
"""把 README 里有、而问卷正文缺失的作答纪律补进 docx 前言页(2026-07-28)。

只动前言页,不碰 gold_sample.json / answer_key / coding_sheet —— 抽样与封存
答案键保持原样。原 docx 备份为 *_原始备份.docx。

缺口:README 写了"不用 AI 辅助",但编码员实际拿到的问卷里没有这句;问卷也
没有禁止在 Stack Overflow 上回搜原题(原页面会同时暴露发布日期和已有答案,
比 AI 辅助更直接地破坏盲态)。

用法: py -V:3.13 patch_questionnaire_instructions.py
"""
import copy
import os
import shutil

import docx
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "人工金标准_问卷_300题.docx")
BAK = os.path.join(HERE, "人工金标准_问卷_300题_原始备份.docx")

RULES = [
    "不得使用任何 AI 工具(ChatGPT / Claude / Copilot / 搜索引擎的 AI 摘要等)"
    "辅助判分。本问卷的全部意义就是取得一份不依赖模型家族的人类基准,"
    "一旦借助模型作答,这份金标准即告失效。",
    "不得在 Stack Overflow 或搜索引擎上回搜原题。原页面会同时暴露发布日期与"
    "已有答案,两者都直接破坏盲态。",
    "凭题面(标题 + 标签 + 正文节选)判分即可。[CODE] 是被折叠的代码块标记,"
    "遇到它就按“这类问题一般需要多少提问者自身的上下文”来判断,不必设法还原代码。",
    "只填 0-4 的整数,不填小数、不留空。确实拿不准时按第一印象给最接近的一档,"
    "并在 coding_sheet 的备注列写一句理由或标“存疑”;不要为了追求前后一致而"
    "回头成批改分。",
    "建议一次连续做 60-100 题,中途可休息;全部 300 题约 2.5-4 小时。",
]

HEAD = "作答纪律(与盲态直接相关,请严格遵守)"


def main():
    if not os.path.exists(BAK):
        shutil.copy2(SRC, BAK)
        print(f"[backup] {BAK}")

    doc = docx.Document(SRC)
    if any(HEAD in p.text for p in doc.paragraphs):
        print("[skip] 纪律段已存在,未重复写入")
        return

    anchor = next(p for p in doc.paragraphs if p.text.startswith("评分录入"))

    head = doc.add_paragraph()
    r = head.add_run(HEAD + ":")
    r.bold = True
    r.font.size = Pt(11)
    anchor._p.addnext(head._p)

    prev = head
    for rule in RULES:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(rule).font.size = Pt(10)
        prev._p.addnext(p._p)
        prev = p

    doc.save(SRC)
    print(f"[saved] {SRC}  (+{len(RULES) + 1} 段)")


if __name__ == "__main__":
    main()
