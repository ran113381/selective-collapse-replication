# -*- coding: utf-8 -*-
"""P9 腿B 人工金标准启动包生成器。

从 60 个月 × 100 题的 python 批次中每月抽 5 题(共 300),盲态(无日期、无模型标签)
随机排序,生成:
  - 人工金标准_问卷_300题.docx      (双人独立编码用;含 rubric 指导语)
  - coding_sheet_A.csv / _B.csv     (数字化评分模板,两位编码员各一份)
  - answer_key_封存勿开.csv          (question_id, 月份, 模型分;编码完成前勿开)
  - gold_sample.json                (抽样存档,含全部呈现字段)
  - score_gold.py                   (κ 计算脚本:人-人 & 人-模型)
  - README.md                       (流程说明)
seed=20260704,全流程确定性。
"""
import json, os, csv, random, re

DATA = r"C:\Users\<user>\AppData\Local\Temp\claude\E-------\77fae8c6-f490-4ed9-96cb-fd023fd3807e\scratchpad\P9\P9_验证经济学_迁移包_20260629\P9_code\legB\data"
OUT = r"E:\智能体论文\P9_金标准_20260704"
os.makedirs(OUT, exist_ok=True)
SEED = 20260704
PER_MONTH = 5

# python 全期 = 原始 37 个月(so_questions_full)+ 扩展 23 个月(so_questions_ext_py)
questions = (json.load(open(os.path.join(DATA, "so_questions_full.json"), encoding="utf-8"))
             + json.load(open(os.path.join(DATA, "so_questions_ext_py.json"), encoding="utf-8")))
labels = {}
for fn in ("question_labels.csv", "question_labels_ext_py.csv"):
    for r in csv.DictReader(open(os.path.join(DATA, fn), encoding="utf-8")):
        labels[int(r["question_id"])] = {"score": int(r["score"]), "label": r["label"]}

by_month = {}
for q in questions:
    by_month.setdefault(q["ym"], []).append(q)
months = sorted(by_month)
assert len(months) == 60, len(months)

rng = random.Random(SEED)
sample = []
for ym in months:
    picks = rng.sample(by_month[ym], PER_MONTH)
    for q in picks:
        lab = labels.get(q["question_id"], {})
        sample.append({
            "question_id": q["question_id"], "ym": ym,
            "title": q["title"], "tags": q["tags"],
            "body_excerpt": q.get("body_excerpt", ""),
            "model_score": lab.get("score"), "model_label": lab.get("label"),
        })
assert all(s["model_score"] is not None for s in sample), "missing labels"

rng.shuffle(sample)              # 盲态呈现顺序
for i, s in enumerate(sample, 1):
    s["order"] = i

json.dump(sample, open(os.path.join(OUT, "gold_sample.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)

# ---- answer key (sealed) ----
with open(os.path.join(OUT, "answer_key_封存勿开.csv"), "w", newline="",
          encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["order", "question_id", "ym", "model_score", "model_label"])
    for s in sample:
        w.writerow([s["order"], s["question_id"], s["ym"],
                    s["model_score"], s["model_label"]])

# ---- coder sheets ----
for coder in ("A", "B"):
    with open(os.path.join(OUT, f"coding_sheet_{coder}.csv"), "w", newline="",
              encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["order", "question_id", "score_0_4", "备注(可空)"])
        for s in sample:
            w.writerow([s["order"], s["question_id"], "", ""])

# ---- questionnaire docx ----
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(10.5)

h = doc.add_heading("P9 腿B 人工金标准编码问卷(300 题,盲态)", level=0)
p = doc.add_paragraph()
p.add_run("用途:为 Stack Overflow 问题的 AI 可替代性分类建立人工金标准,"
          "校验 LLM 分类器(独立于模型家族)。两位编码员独立作答,期间不得讨论、"
          "不得查看对方评分、不得查询问题的发布日期。").bold = True
doc.add_paragraph("评分录入:请在 coding_sheet_A.csv / coding_sheet_B.csv 的 score_0_4 "
                  "列填写 0-4 整数(与题目 order 对应);本问卷仅作阅读材料。")

doc.add_heading("评分标准(0-4 AI 可替代性)", level=1)
rubric = [
    ("4 = 纯 GENERATION", "概念/算法/标准 how-to,自足,有唯一标准答案,LLM 一次答对。"),
    ("3 = 偏 generation", "标准做法 + 轻微上下文。"),
    ("2 = 混合/偏 verification", "需要一些判断或提问者的具体设置。"),
    ("1 = VERIFICATION", "调试提问者的具体报错/代码/环境、在其数据上调性能、生产环境特定。"),
    ("0 = 纯 verification", "完全受上下文/判断约束,仅凭文本无法回答。"),
]
for k, v in rubric:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(k + ":").bold = True
    p.add_run(v)
doc.add_paragraph("关键原则:按问题的内在类型判定;判断核心是——答案是“生成一段标准代码/"
                  "解释”(高分)还是“对这个人的具体情况做诊断/权衡”(低分)。"
                  "[CODE] 为被折叠的代码块标记。")
doc.add_page_break()

def clean(t):
    t = re.sub(r"\s+", " ", t or "").strip()
    return t[:1200] + (" …[截断]" if len(t) > 1200 else "")

for s in sample:
    p = doc.add_paragraph()
    r = p.add_run(f"[{s['order']:03d}]  {s['title']}")
    r.bold = True; r.font.size = Pt(11)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("tags: " + ", ".join(s["tags"]))
    r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    doc.add_paragraph(clean(s["body_excerpt"]))
    p3 = doc.add_paragraph()
    r3 = p3.add_run("评分(0-4):____")
    r3.font.color.rgb = RGBColor(0x2E, 0x5A, 0x88)

doc.save(os.path.join(OUT, "人工金标准_问卷_300题.docx"))
print("docx saved,", len(sample), "items")

# ---- score_gold.py ----
SCORE = '''# -*- coding: utf-8 -*-
"""金标准 κ 计算:人-人 与 人-模型(二元 GEN/VER + 二次加权 0-4)。
用法: python score_gold.py   (需 coding_sheet_A/B.csv 已填完)"""
import csv

def load(path, col):
    d = {}
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        v = r[col].strip()
        if v != "":
            d[int(r["order"])] = int(v)
    return d

def kappa_binary(x, y, cut=3):
    ks = sorted(set(x) & set(y))
    a = [int(x[k] >= cut) for k in ks]; b = [int(y[k] >= cut) for k in ks]
    n = len(ks); po = sum(i == j for i, j in zip(a, b)) / n
    pa1 = sum(a) / n; pb1 = sum(b) / n
    pe = pa1 * pb1 + (1 - pa1) * (1 - pb1)
    return (po - pe) / (1 - pe), po, n

def kappa_qw(x, y, K=5):
    ks = sorted(set(x) & set(y)); n = len(ks)
    O = [[0] * K for _ in range(K)]
    for k in ks:
        O[x[k]][y[k]] += 1
    rx = [sum(O[i]) for i in range(K)]; ry = [sum(O[i][j] for i in range(K)) for j in range(K)]
    num = den = 0.0
    for i in range(K):
        for j in range(K):
            w = ((i - j) ** 2) / ((K - 1) ** 2)
            num += w * O[i][j]
            den += w * rx[i] * ry[j] / n
    return 1 - num / den, n

A = load("coding_sheet_A.csv", "score_0_4")
B = load("coding_sheet_B.csv", "score_0_4")
M = load("answer_key_封存勿开.csv", "model_score")
for name, x, y in [("人A-人B", A, B), ("人A-模型", A, M), ("人B-模型", B, M)]:
    kb, po, n = kappa_binary(x, y)
    kw, _ = kappa_qw(x, y)
    print(f"{name}: N={n}  二元κ={kb:.3f} (一致率{po:.1%})  二次加权κ={kw:.3f}")
'''
open(os.path.join(OUT, "score_gold.py"), "w", encoding="utf-8").write(SCORE)

# ---- README ----
README = """# P9 腿B 人工金标准启动包(2026-07-04)

**目的**:用两位人类编码员为 300 个 Stack Overflow python 问题建立 AI 可替代性
金标准,消解"两位 LLM 评分员同属 Claude 家族"这一分类器效度攻击点。

**抽样**:每月 5 题 × 60 月(2021-06 至 2026-05),seed=20260704,分层保证时间覆盖;
呈现顺序整体随机打乱,题面不含日期(盲态,防止编码员推断 ChatGPT 前后)。

**流程**:
1. 两位编码员各持《人工金标准_问卷_300题.docx》独立阅读,在各自的
   coding_sheet_A/B.csv 填 0-4 分;期间不讨论、不查发布日期、不用 AI 辅助。
2. 全部填完前,任何人不得打开 answer_key_封存勿开.csv。
3. 完成后运行 `python score_gold.py`,输出人-人与人-模型的二元 κ 与二次加权 κ。
4. 论文引用口径:人-人 κ 为金标准可靠性;人-模型 κ 为分类器效度;
   与既有同族 κ(0.65/0.78)并列报告。

**预期工时**:约 2.5-4 小时/人(300 题 × 30-45 秒)。
**生成脚本**:make_gold_package.py(replication_additions 亦有存档);全流程确定性可复现。
"""
open(os.path.join(OUT, "README.md"), "w", encoding="utf-8").write(README)
print("package complete ->", OUT)
